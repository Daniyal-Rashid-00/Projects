import os
import re
from bs4 import BeautifulSoup
from docx import Document

# Path to the folder containing HTML files
folder_path = r"D:\Keep"

# Output folder for DOCX files
output_folder = os.path.join(folder_path, "DOCX_Files")
os.makedirs(output_folder, exist_ok=True)

# Regex pattern to detect timestamps (e.g., "Aug 4, 2024, 6:10:10 PM")
timestamp_pattern = re.compile(r"\b[A-Za-z]{3,9} \d{1,2}, \d{4}, \d{1,2}:\d{1,2}(:\d{1,2})? [APap][Mm]\b")

# Regex to remove invalid filename characters
invalid_chars = re.compile(r'[\/:*?"<>|]')

for filename in os.listdir(folder_path):
    if filename.endswith(".html"):
        file_path = os.path.join(folder_path, filename)

        with open(file_path, "r", encoding="utf-8") as file:
            html_content = file.read()

            # Parse HTML content with BeautifulSoup
            soup = BeautifulSoup(html_content, "html.parser")

            # Extract text and remove unwanted timestamps
            extracted_lines = []
            for element in soup.stripped_strings:
                if not timestamp_pattern.match(element):  # Skip timestamps
                    extracted_lines.append(element)

            # Remove duplicate links (only keep the first occurrence)
            seen_links = set()
            cleaned_lines = []
            for line in extracted_lines:
                if line.startswith("http"):  # If it's a link
                    if line in seen_links:  # Skip duplicate links
                        continue
                    seen_links.add(line)
                cleaned_lines.append(line)

            # Get the third non-empty line as the title
            title = (
                cleaned_lines[2].replace("*", "").strip()
                if len(cleaned_lines) >= 3
                else filename.replace(".html", "")  # Use filename if less than 3 lines exist
            )
            title = re.sub(invalid_chars, "", title)  # Remove invalid filename characters

            # Create a new DOCX document
            doc = Document()

            for line in cleaned_lines:
                cleaned_line = line.replace("*", "").strip()  # Remove '*' but keep text

                if cleaned_line.startswith("★"):  # Add bullet points
                    doc.add_paragraph(cleaned_line, style="List Bullet")
                elif line.startswith("*"):  # Convert bold headings
                    doc.add_paragraph(cleaned_line, style="Heading 2")
                else:
                    doc.add_paragraph(cleaned_line)

            # Save the cleaned and formatted DOCX file with the third line as the filename
            output_file = os.path.join(output_folder, f"{title}.docx")
            doc.save(output_file)

print(f"Conversion complete! Check '{output_folder}' for the renamed DOCX files.")
